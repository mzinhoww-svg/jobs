import tempfile
import unittest
from pathlib import Path
from docx import Document
from pypdf import PdfReader
from jobs_service.documents import generate_documents


class DocumentTests(unittest.TestCase):
    def test_nonempty_specific_documents_parse_and_hash(self):
        p={'full_name':'Fictional Candidate','preferred_email':'candidate@example.invalid','location':'São Paulo, Brazil','facts':[{'id':'f1','text':'Structured payment partnerships across product and operations.','source':'controlled fixture','date':'2026-10-27','source_kind':'candidate_declaration','confirmation':'confirmed'}]}
        job={'title':'Payments Product Lead','company':'Example Industries','country':'BR','language':'en','description':'Develop payment partnerships and coordinate product teams.'}
        with tempfile.TemporaryDirectory() as d:
            docs=generate_documents(p,job,Path(d))
            for name in ('cv.pdf','cv.docx','cover-letter.pdf','cover-letter.docx'):
                self.assertGreater((Path(d)/name).stat().st_size,1000)
                self.assertEqual(len(docs[name]['sha256']),64)
            text='\n'.join(page.extract_text() for page in PdfReader(Path(d)/'cv.pdf').pages)
            self.assertIn('Payments Product Lead',text)
            self.assertIn('São Paulo',text)
            self.assertNotIn('Miami',text)
            self.assertIn('Structured payment partnerships',text)
            fonts=PdfReader(Path(d)/'cv.pdf').pages[0]['/Resources']['/Font'].get_object()
            embedded=[f.get_object().get('/FontDescriptor') for f in fonts.values()]
            self.assertTrue(any(f and '/FontFile2' in f.get_object() for f in embedded),'PDF must embed a readable portable font')
            doc=Document(Path(d)/'cv.docx')
            self.assertEqual(len(doc.tables),0)
            self.assertIn('Fictional Candidate',doc.paragraphs[0].text)

    def test_unconfirmed_claim_excluded(self):
        p={'full_name':'Example','preferred_email':'a@example.invalid','location':'Miami, FL, USA','facts':[{'id':'x','text':'Invented achievement','confirmation':'reference'}]}
        with tempfile.TemporaryDirectory() as d:
            generate_documents(p,{'title':'Lead','company':'Example','description':'Lead team','language':'en'},Path(d))
            self.assertNotIn('Invented', '\n'.join(x.text for x in Document(Path(d)/'cv.docx').paragraphs))

    def test_translated_evidence_uses_job_language_and_keeps_fact_identity(self):
        import json
        profile={'full_name':'Fictional Candidate','preferred_email':'fixture@example.invalid','location':'São Paulo, Brazil','facts':[{'id':'translation-1','text':'Novos negócios e parcerias','translations':{'en':'New business development and partnerships'},'source':'controlled fixture','date':'2026-10-27','source_kind':'candidate_declaration','confirmation':'confirmed'}]}
        with tempfile.TemporaryDirectory() as tmp:
            generate_documents(profile,{'title':'Partnerships Lead','company':'Fictional Labs','language':'en','description':'Develop partnerships and lead new business opportunities.'},Path(tmp))
            text=(Path(tmp)/'cv.txt').read_text()
            self.assertIn('New business development and partnerships',text)
            self.assertNotIn('Novos negócios',text)
            report=json.loads((Path(tmp)/'tailoring-report.json').read_text())
            self.assertEqual(report['selected_fact_ids'],['translation-1'])
            self.assertEqual(report['method'],'lexical_evidence_matching')
